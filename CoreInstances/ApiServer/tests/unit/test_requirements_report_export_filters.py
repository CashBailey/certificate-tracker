from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from unittest.mock import AsyncMock

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.deps import get_clock, get_current_user, get_repository
from src.routes.reports import router as reports_router
from src.shared.models import CertificateType, Employee, RequirementAssignment, RequirementStatus, Role


def _employee(employee_id: int, first_name: str, last_name: str, role: Role = Role.EMPLOYEE) -> Employee:
    now = datetime.now(timezone.utc)
    return Employee(
        id=employee_id,
        employee_number=f"E{employee_id:05d}",
        first_name=first_name,
        last_name=last_name,
        email=f"{first_name.lower()}.{last_name.lower()}@ci.laredo.tx.us",
        role=role,
        is_active=True,
        created_at=now,
        updated_at=now,
    )


def _requirement(
    requirement_id: int,
    employee_id: int,
    certificate_type_id: int,
    due_date: date,
    *,
    status: RequirementStatus = RequirementStatus.NOT_STARTED,
    satisfied_by_id: int | None = None,
    waived_at=None,
) -> RequirementAssignment:
    now = datetime.now(timezone.utc)
    return RequirementAssignment(
        id=requirement_id,
        employee_id=employee_id,
        certificate_type_id=certificate_type_id,
        due_date=due_date,
        status=status,
        satisfied_by_id=satisfied_by_id,
        waived_at=waived_at,
        created_at=now,
        updated_at=now,
    )


class _FixedClock:
    def __init__(self, today_value: date):
        self._today = today_value

    def now(self) -> datetime:
        return datetime.combine(self._today, datetime.min.time(), tzinfo=timezone.utc)

    def today(self) -> date:
        return self._today


def _build_client() -> TestClient:
    today = date(2026, 4, 20)
    coordinator = _employee(9000, "Casey", "Coordinator", role=Role.COORDINATOR)
    rachel = _employee(1001, "Rachel", "Sullivan")
    bob = _employee(1002, "Bob", "Builder")

    cert_types = [
        CertificateType(id=1, name="Forklift", description="Forklift cert"),
        CertificateType(id=2, name="Hazmat", description="Hazmat cert"),
    ]

    requirements = [
        _requirement(1, rachel.id, 1, today - timedelta(days=1)),
        _requirement(2, rachel.id, 2, today + timedelta(days=5)),
        _requirement(3, bob.id, 1, today + timedelta(days=60), status=RequirementStatus.IN_PROGRESS),
        _requirement(
            4,
            bob.id,
            2,
            today + timedelta(days=10),
            status=RequirementStatus.SATISFIED,
            satisfied_by_id=77,
        ),
    ]

    repo = AsyncMock()
    repo.list_employees = AsyncMock(return_value=[rachel, bob])
    repo.list_certificate_types = AsyncMock(return_value=cert_types)
    repo.list_all_requirements = AsyncMock(return_value=requirements)

    app = FastAPI()
    app.include_router(reports_router, prefix="/api")
    app.dependency_overrides[get_repository] = lambda: repo
    app.dependency_overrides[get_current_user] = lambda: coordinator
    app.dependency_overrides[get_clock] = lambda: _FixedClock(today)
    return TestClient(app)


def test_requirements_csv_export_respects_search_and_requirement_status_filters():
    client = _build_client()

    response = client.get(
        "/api/reports/requirements?format=csv&requirement_status=Overdue&search=Rachel"
    )

    assert response.status_code == 200
    lines = [line for line in response.text.strip().splitlines() if line]
    assert len(lines) == 2
    assert "Rachel Sullivan" in lines[1]
    assert "Overdue" in lines[1]
    assert "Bob Builder" not in response.text
    assert "DueSoon" not in response.text


def test_requirements_csv_export_satisfied_filter_excludes_on_track_rows():
    client = _build_client()

    response = client.get("/api/reports/requirements?format=csv&requirement_status=Satisfied")

    assert response.status_code == 200
    assert "Bob Builder" in response.text
    assert "Satisfied" in response.text
    assert "InProgress" not in response.text
    assert response.text.count("Bob Builder") == 1


def test_requirements_xlsx_export_summary_uses_filtered_row_count():
    client = _build_client()

    response = client.get("/api/reports/requirements?format=xlsx&requirement_status=InProgress")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    summary = workbook["Summary"]
    detail = workbook["Requirements"]

    summary_metrics = {summary.cell(row=row, column=1).value: summary.cell(row=row, column=2).value for row in range(2, 9)}
    assert summary_metrics["Total Requirements"] == 1
    assert summary_metrics["Compliant"] == 1
    assert detail.cell(row=2, column=1).value == "Bob Builder"
    assert detail.cell(row=2, column=5).value == "InProgress"
    assert detail.max_row == 2


def test_requirements_export_rejects_invalid_requirement_status():
    client = _build_client()

    response = client.get("/api/reports/requirements?format=csv&requirement_status=Nope")

    assert response.status_code == 400
    assert "Invalid requirement_status" in response.text
