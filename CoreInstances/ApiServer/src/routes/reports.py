"""
Reports API endpoints.
"""

import csv
import io
from collections import defaultdict
from datetime import date
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from ..deps import get_repository, get_clock, require_coordinator
from ..shared.models import Employee
from ..shared.protocols import Clock, Repository
from ..shared.utils import (
    compute_requirement_compliance_status,
)
from .schemas import (
    ComplianceReportResponse,
    ComplianceStatusCount,
)

router = APIRouter(prefix="/reports", tags=["reports"])


def _derive_requirements_table_status(req, current_date: date) -> str:
    """Mirror the Requirements page status buckets used for table filtering/export."""
    if req.waived_at:
        return "Waived"
    if req.satisfied_by_id:
        return "Satisfied"
    if req.due_date < current_date:
        return "Overdue"
    days_until_due = (req.due_date - current_date).days
    if days_until_due <= 30:
        return "DueSoon"
    return "InProgress"


def _count_compliance_statuses(rows: list[dict]) -> ComplianceStatusCount:
    counts = ComplianceStatusCount()
    for row in rows:
        compliance_status = row["compliance_status"]
        if compliance_status == "Compliant":
            counts.compliant += 1
        elif compliance_status == "DueSoon":
            counts.due_soon += 1
        elif compliance_status == "Overdue":
            counts.overdue += 1
        elif compliance_status == "Waived":
            counts.waived += 1
    return counts


@router.get("/requirements")
async def get_requirements_compliance_report(
    format: Optional[str] = Query(
        "json",
        description="Output format: 'json', 'csv', or 'xlsx'",
        pattern=r"^(json|csv|xlsx)$",
    ),
    status: Optional[str] = Query(
        None,
        description="Filter by compliance status: Compliant, DueSoon, Overdue, Waived",
    ),
    requirement_status: Optional[str] = Query(
        None,
        description="Filter exported rows by requirements table status: Overdue, DueSoon, InProgress, Satisfied, Waived",
    ),
    search: Optional[str] = Query(
        None,
        description="Filter exported rows by employee full name or certificate type",
    ),
    current_user: Employee = Depends(require_coordinator),
    repository: Repository = Depends(get_repository),
    clock: Clock = Depends(get_clock),
):
    """
    Get requirements compliance dashboard data.

    Only Coordinator can view compliance reports.
    Returns JSON by default, or XLSX file when format=xlsx.
    """
    current_date = clock.today()

    # Bulk-load all data in 3 queries (avoids N+1 pattern). Pass limit=None
    # so compliance counters and CSV/XLSX exports reflect every row, not a
    # silently-truncated 10k-row sample.
    employees = await repository.list_employees()
    cert_types = await repository.list_certificate_types()
    all_requirements = await repository.list_all_requirements(limit=None)

    # Filter out Admin system accounts from employee count (B1/B2)
    employees = [e for e in employees if e.role != "Admin"]
    total_employees = len(employees)

    # Pre-build lookup maps
    employee_map = {e.id: e for e in employees}
    employee_ids = set(employee_map.keys())
    cert_type_map = {ct.id: ct for ct in cert_types}

    # Group requirements by employee (only non-Admin employees)
    reqs_by_employee: dict[int, list] = defaultdict(list)
    for req in all_requirements:
        if req.employee_id in employee_ids:
            reqs_by_employee[req.employee_id].append(req)

    total_requirements = 0
    overall_counts = ComplianceStatusCount()
    by_cert_type: dict[str, ComplianceStatusCount] = {
        ct.name: ComplianceStatusCount() for ct in cert_types
    }
    requirements_rows: list[dict] = []

    # Process — no DB calls in this loop
    for employee in employees:
        for req in reqs_by_employee.get(employee.id, []):
            total_requirements += 1

            table_status = _derive_requirements_table_status(req, current_date)
            compliance_status = compute_requirement_compliance_status(
                req.status, req.due_date, current_date,
            )

            # Update overall counts
            if compliance_status.value == "Compliant":
                overall_counts.compliant += 1
            elif compliance_status.value == "DueSoon":
                overall_counts.due_soon += 1
            elif compliance_status.value == "Overdue":
                overall_counts.overdue += 1
            elif compliance_status.value == "Waived":
                overall_counts.waived += 1

            # Update per-type counts using pre-built map
            cert_type = cert_type_map.get(req.certificate_type_id)
            if cert_type and cert_type.name in by_cert_type:
                type_counts = by_cert_type[cert_type.name]
                if compliance_status.value == "Compliant":
                    type_counts.compliant += 1
                elif compliance_status.value == "DueSoon":
                    type_counts.due_soon += 1
                elif compliance_status.value == "Overdue":
                    type_counts.overdue += 1
                elif compliance_status.value == "Waived":
                    type_counts.waived += 1

            requirements_rows.append({
                "employee_name": f"{employee.first_name} {employee.last_name}".strip(),
                "employee_number": employee.employee_number,
                "certificate_type": cert_type.name if cert_type else "Unknown",
                "due_date": req.due_date,
                "status": table_status,
                "compliance_status": compliance_status.value,
            })

    if format in ("xlsx", "csv"):
        # Filter rows by compliance status if requested. Same filter logic for
        # both file formats — keep them in sync so a Coordinator exporting CSV
        # vs XLSX with the same status filter gets the same row set.
        rows_for_export = requirements_rows
        if status:
            valid_statuses = {"Compliant", "DueSoon", "Overdue", "Waived"}
            if status not in valid_statuses:
                raise HTTPException(
                    status_code=400,
                    detail=f"Invalid status: {status}. Must be one of: {', '.join(sorted(valid_statuses))}",
                )
            rows_for_export = [r for r in requirements_rows if r["compliance_status"] == status]
        if requirement_status:
            valid_requirement_statuses = {"Overdue", "DueSoon", "InProgress", "Satisfied", "Waived"}
            if requirement_status not in valid_requirement_statuses:
                raise HTTPException(
                    status_code=400,
                    detail="Invalid requirement_status: "
                    f"{requirement_status}. Must be one of: {', '.join(sorted(valid_requirement_statuses))}",
                )
            rows_for_export = [r for r in rows_for_export if r["status"] == requirement_status]
        if search:
            search_term = search.strip().lower()
            if search_term:
                rows_for_export = [
                    r
                    for r in rows_for_export
                    if search_term in r["employee_name"].lower()
                    or search_term in r["certificate_type"].lower()
                ]

        counts_for_export = _count_compliance_statuses(rows_for_export)
        export_is_filtered = bool(status or requirement_status or (search and search.strip()))

        if format == "xlsx":
            return _generate_requirements_xlsx_response(
                rows_for_export, counts_for_export,
                total_employees,
                len(rows_for_export) if export_is_filtered else total_requirements,
                current_date,
            )
        return _generate_requirements_csv_response(rows_for_export, current_date)

    return ComplianceReportResponse(
        total_employees=total_employees,
        total_requirements=total_requirements,
        status_counts=overall_counts,
        by_certificate_type=by_cert_type,
    )


def _safe_csv_cell(value) -> str:
    """Defuse CSV formula-injection attempts (MED-07).

    Python's `csv.writer` handles RFC 4180 quoting (commas, quotes, newlines)
    correctly, but it does NOT defang values that begin with =, +, -, @, or
    leading whitespace control chars — Excel/LibreOffice/Sheets interpret
    those as formulas when the cell is opened. An attacker-controlled field
    like `=CMD|'/c calc'!A0` would execute when a Coordinator opens the
    export.

    Defense: prepend a single quote (`'`) to any cell whose first character
    is one of the dangerous prefixes. Excel hides the leading quote as a
    text-entry marker; the formula is rendered as inert text. This is the
    standard OWASP CSV-injection mitigation. Apply at every CSV writer
    site that emits user-controlled data.
    """
    s = "" if value is None else str(value)
    if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + s
    return s


def _safe_xlsx_cell(value) -> str:
    """Force attacker-controlled spreadsheet values to render as text."""
    return _safe_csv_cell(value)


def _generate_requirements_csv_response(
    rows: list[dict], current_date: date
) -> StreamingResponse:
    """Generate CSV file response for requirements compliance report.

    Mirrors the column set used in the XLSX export so a Coordinator
    exporting the same filter as XLSX vs CSV gets identical content
    in different containers. All user-controllable fields run through
    _safe_csv_cell() to defuse formula-injection attacks (MED-07).
    """
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "Employee Number",
            "Employee Name",
            "Certificate Type",
            "Due Date",
            "Requirement Status",
            "Compliance Status",
        ]
    )
    for row in rows:
        due = row["due_date"]
        writer.writerow(
            [
                _safe_csv_cell(row["employee_number"]),
                _safe_csv_cell(row["employee_name"]),
                _safe_csv_cell(row["certificate_type"]),
                _safe_csv_cell(due.isoformat() if due else ""),
                _safe_csv_cell(row["status"]),
                _safe_csv_cell(row["compliance_status"]),
            ]
        )
    output.seek(0)
    filename = f"requirements_compliance_{current_date.isoformat()}.csv"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ==================== XLSX Generation ====================

_XLSX_HEADER_FONT = Font(bold=True, color="FFFFFF")
_XLSX_HEADER_FILL = PatternFill(
    start_color="4472C4", end_color="4472C4", fill_type="solid"
)


def _write_xlsx_header(ws, headers: list[str]) -> None:
    """Write styled header row to a worksheet."""
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _XLSX_HEADER_FONT
        cell.fill = _XLSX_HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _auto_width_columns(ws) -> None:
    """Auto-size column widths based on content."""
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 2, 40)


def _generate_requirements_xlsx_response(
    requirements_rows: list[dict],
    overall_counts: ComplianceStatusCount,
    total_employees: int,
    total_requirements: int,
    report_date: date,
) -> StreamingResponse:
    """Generate XLSX for requirements compliance report with Summary + Detail sheets."""
    wb = Workbook()

    # Summary sheet
    ws_summary = wb.active
    ws_summary.title = "Summary"
    _write_xlsx_header(ws_summary, ["Metric", "Value"])
    summary_data = [
        ("Report Date", report_date.isoformat()),
        ("Total Employees", total_employees),
        ("Total Requirements", total_requirements),
        ("Compliant", overall_counts.compliant),
        ("Due Soon", overall_counts.due_soon),
        ("Overdue", overall_counts.overdue),
        ("Waived", overall_counts.waived),
    ]
    for row_idx, (metric, value) in enumerate(summary_data, 2):
        ws_summary.cell(row=row_idx, column=1, value=metric)
        ws_summary.cell(row=row_idx, column=2, value=value)
    _auto_width_columns(ws_summary)

    # Detail sheet
    ws_detail = wb.create_sheet("Requirements")
    detail_headers = [
        "Employee", "Employee Number", "Certificate Type",
        "Due Date", "Status", "Compliance Status",
    ]
    _write_xlsx_header(ws_detail, detail_headers)

    for row_idx, row_data in enumerate(requirements_rows, 2):
        ws_detail.cell(row=row_idx, column=1, value=_safe_xlsx_cell(row_data["employee_name"]))
        ws_detail.cell(row=row_idx, column=2, value=_safe_xlsx_cell(row_data["employee_number"]))
        ws_detail.cell(row=row_idx, column=3, value=_safe_xlsx_cell(row_data["certificate_type"]))
        ws_detail.cell(row=row_idx, column=4, value=row_data["due_date"].isoformat() if row_data["due_date"] else "")
        ws_detail.cell(row=row_idx, column=5, value=_safe_xlsx_cell(row_data["status"]))
        ws_detail.cell(row=row_idx, column=6, value=_safe_xlsx_cell(row_data["compliance_status"]))
    _auto_width_columns(ws_detail)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="requirements_compliance_{report_date.isoformat()}.xlsx"',
        },
    )
